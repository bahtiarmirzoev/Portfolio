import logging
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional, Tuple

import fitz  # PyMuPDF
from docx import Document
from openpyxl import load_workbook

from config import INPUT_PDF_DIR, LOG_FILE, OUTPUT_DIR, TEMPLATES_DIR


class ProcessorError(Exception):
    """Base processor exception."""


class DataNotRecognizedError(ProcessorError):
    """Raised when required data cannot be extracted."""


@dataclass
class ExtractedData:
    court: str
    case_no: str
    car_model: str
    car_colour: str
    plate_no: str
    judge: str
    judge_short: str
    accident_date: str
    total_amount: str
    total_amount_str: str
    report_date: str


MONTH_NAMES_AZ = (
    "yanvar|fevral|mart|aprel|may|iyun|iyul|avqust|sentyabr|oktyabr|noyabr|dekabr"
)
DATE_TOKEN_PATTERN = (
    rf"(?:\d{{2}}[./-]\d{{2}}[./-]\d{{4}}|\d{{1,2}}\s+(?:{MONTH_NAMES_AZ})\s+\d{{4}}(?:-cü|-ci|-cu|-cı)?(?:\s+il)?)"
)
DEFAULT_CAR_COLOUR = ""


def ensure_directories() -> None:
    INPUT_PDF_DIR.mkdir(parents=True, exist_ok=True)
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def setup_logging() -> None:
    logging.basicConfig(
        filename=LOG_FILE,
        level=logging.INFO,
        filemode="w",
        format="%(asctime)s | %(levelname)s | %(message)s",
        encoding="utf-8",
    )


def find_first_file(directory: Path, patterns: tuple[str, ...]) -> Optional[Path]:
    for pattern in patterns:
        files = sorted(directory.glob(pattern))
        if files:
            return files[0]
    return None


def read_pdf_text(pdf_path: Path) -> str:
    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF file not found: {pdf_path}")

    full_text: list[str] = []
    with fitz.open(pdf_path) as doc:
        for page in doc:
            full_text.append(page.get_text("text"))
    return "\n".join(full_text)


def extract_with_regex(text: str, pattern: str, flags: int = 0) -> Optional[str]:
    match = re.search(pattern, text, flags)
    if not match:
        return None
    return match.group(1).strip()


def extract_court_name(text: str) -> Optional[str]:
    first_page_fragment = "\n".join(text.splitlines()[:40])
    specific = extract_with_regex(
        first_page_fragment,
        r"(Bakı\s+Apellyasiya\s+Məhkəməsi)",
        flags=re.IGNORECASE,
    )
    if specific:
        return specific

    lines = [line.strip() for line in first_page_fragment.splitlines() if line.strip()]
    for line in lines:
        line = re.sub(r"\s+", " ", line)
        match = re.search(
            r"((?:Bakı(?:\s+Şəhəri)?|Sumqayıt|Gəncə|Naxçıvan|Şəki|Şirvan|Lənkəran|Mingəçevir|Yevlax|Quba|Xaçmaz|Şamaxı|Qəbələ|Ağdaş|Sabirabad|Salyan|Masallı|Cəlilabad|Abşeron|Xırdalan|Qazax|Tovuz|Ağstafa|Zaqatala|Balakən|Qax|Bərdə|Ağcabədi|Ağdam|Füzuli|Beyləqan|İmişli|Saatlı|Kürdəmir|Ucar|Göyçay|Zərdab|Goranboy|Naftalan|Samux|Daşkəsən|Gədəbəy|Şəmkir|Qusar|Siyəzən|Şabran|Neftçala|Biləsuvar|Yardımlı|Lerik|Astara)[^\n,;:]{0,140}?(?:Rayon Məhkəməsi|Apellyasiya Məhkəməsi))",
            line,
            flags=re.IGNORECASE,
        )
        if match:
            return match.group(1).strip()

    return extract_with_regex(
        first_page_fragment,
        r"([A-ZƏÖÜĞÇŞİA-Za-zəöüğçşı\s]{3,160}?(?:Rayon Məhkəməsi|Apellyasiya Məhkəməsi))",
        flags=re.IGNORECASE,
    )


def extract_case_no(text: str) -> Optional[str]:
    return extract_with_regex(text, r"\b(\d+\([0-9\-]+\)-\d+/\d{4})\b")


def extract_plate_no(text: str) -> Optional[str]:
    return extract_with_regex(text, r"\b(\d{2}-[A-Z]{2}-\d{3})\b")


def extract_car_model(text: str) -> Optional[str]:
    quoted = extract_with_regex(
        text,
        r"[\"“”'«»]\s*([^\"“”'«»\n]{2,100}?)\s*[\"“”'«»]\s*markalı",
        flags=re.IGNORECASE,
    )
    if quoted:
        return quoted

    return extract_with_regex(
        text,
        r"\b([A-Z0-9][A-Za-zƏÖÜĞÇŞİəöüğçşı0-9\-\s]{1,80})\s+markalı",
        flags=re.IGNORECASE,
    )


def extract_car_colour(text: str) -> Optional[str]:
    return extract_with_regex(
        text,
        r"\b([A-Za-zƏÖÜĞÇŞİəöüğçşı]+(?:\s+[A-Za-zƏÖÜĞÇŞİəöüğçşı]+)?)\s+rəngli\b",
        flags=re.IGNORECASE,
    )


def extract_incident_date(text: str) -> Optional[str]:
    return extract_with_regex(
        text,
        rf"({DATE_TOKEN_PATTERN})\s+tarixdə\s+baş\s+vermiş\s+yol-nəqliyyat\s+hadisəsi",
        flags=re.IGNORECASE,
    )


def extract_full_names_from_block(block: str) -> str:
    names = re.findall(
        r"\b([A-ZƏÖÜĞÇŞİ][a-zəöüğçşı]+(?:\s+[A-ZƏÖÜĞÇŞİ][a-zəöüğçşı]+){1,2})\b",
        block,
    )
    seen = set()
    unique = []
    for name in names:
        cleaned = re.sub(r"\s+", " ", name).strip()
        key = cleaned.lower()
        if key not in seen:
            seen.add(key)
            unique.append(cleaned)
    return ", ".join(unique)


def extract_judge(text: str, court: str) -> Optional[str]:
    is_appellate = "apellyasiya məhkəməsi" in court.lower() or bool(
        re.search(r"Apellyasiya\s+Məhkəməsi", text, re.IGNORECASE)
    )
    if is_appellate:
        judges_block = extract_with_regex(
            text,
            r"Hakimlər\s*(.+?)\s*tərkibdə",
            flags=re.IGNORECASE | re.DOTALL,
        )
        if not judges_block:
            return None
        return extract_full_names_from_block(judges_block) or None

    rayon_judge = extract_with_regex(
        text,
        r"hakimi\s*(?:cənab|xanım)\s+([A-ZƏÖÜĞÇŞİA-Za-zəöüğçşı\.\-\s]+?)(?:[,\n]|$)",
        flags=re.IGNORECASE,
    )
    if not rayon_judge:
        return None
    return rayon_judge


def get_judge_short_label(judge: str) -> str:
    if "," in judge:
        return "Hakimlər"
    if judge.strip():
        return "Hakim"
    return ""


def extract_pdf_data(text: str) -> Tuple[Dict[str, str], list[str]]:
    extracted: Dict[str, str] = {}
    missing: list[str] = []

    court = extract_court_name(text)
    if court:
        extracted["court"] = court
    else:
        missing.append("court")

    case_no = extract_case_no(text)
    if case_no:
        extracted["case_no"] = case_no
    else:
        missing.append("case_no")

    judge = extract_judge(text, court or "")
    if judge:
        extracted["judge"] = judge
    else:
        missing.append("judge")

    car_model = extract_car_model(text)
    if car_model:
        extracted["car_model"] = car_model
    else:
        missing.append("car_model")

    car_colour = extract_car_colour(text)
    if car_colour:
        extracted["car_colour"] = car_colour
    else:
        missing.append("car_colour")

    plate_no = extract_plate_no(text)
    if plate_no:
        extracted["plate_no"] = plate_no
    else:
        missing.append("plate_no")

    accident_date = extract_incident_date(text)
    if accident_date:
        extracted["accident_date"] = accident_date
    else:
        missing.append("accident_date")

    return extracted, missing


def get_totals_from_excel(excel_path: Path) -> str:
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    workbook = load_workbook(excel_path, data_only=True)
    try:
        worksheet = workbook.active
        for row_idx in range(1, worksheet.max_row + 1):
            for col_idx in range(1, worksheet.max_column + 1):
                cell_value = worksheet.cell(row=row_idx, column=col_idx).value
                cell_text = str(cell_value).strip() if cell_value is not None else ""
                if cell_text.lower() == "yekun dəyər":
                    d_value = worksheet.cell(row=row_idx, column=4).value
                    e_value = worksheet.cell(row=row_idx, column=5).value

                    if d_value not in (None, ""):
                        return str(d_value).strip()
                    if e_value not in (None, ""):
                        return str(e_value).strip()
                    return ""
    finally:
        workbook.close()

    return ""


def parse_amount_to_int(amount: str) -> Optional[int]:
    if not amount:
        return None
    normalized = str(amount).replace(" ", "").replace(",", ".")
    match = re.search(r"\d+(?:\.\d+)?", normalized)
    if not match:
        return None
    return int(float(match.group(0)))


def number_to_az_words(number: int) -> str:
    units = [
        "sıfır",
        "bir",
        "iki",
        "üç",
        "dörd",
        "beş",
        "altı",
        "yeddi",
        "səkkiz",
        "doqquz",
    ]
    tens = ["", "on", "iyirmi", "otuz", "qırx", "əlli", "altmış", "yetmiş", "səksən", "doxsan"]

    if number == 0:
        return units[0]

    def three_digits_to_words(n: int) -> str:
        parts: list[str] = []
        hundred = n // 100
        rest = n % 100
        ten = rest // 10
        unit = rest % 10

        if hundred:
            if hundred == 1:
                parts.append("yüz")
            else:
                parts.append(f"{units[hundred]} yüz")
        if ten:
            parts.append(tens[ten])
        if unit:
            parts.append(units[unit])
        return " ".join(parts)

    parts: list[str] = []
    millions = number // 1_000_000
    thousands = (number % 1_000_000) // 1_000
    remainder = number % 1_000

    if millions:
        parts.append(f"{three_digits_to_words(millions)} milyon")
    if thousands:
        if thousands == 1:
            parts.append("min")
        else:
            parts.append(f"{three_digits_to_words(thousands)} min")
    if remainder:
        parts.append(three_digits_to_words(remainder))

    return " ".join(part for part in parts if part).strip()


def amount_to_az_text(amount: str) -> str:
    number_value = parse_amount_to_int(amount)
    if number_value is None:
        return ""
    return number_to_az_words(number_value)


def extract_total_amount_from_excel(excel_path: Path) -> str:
    return get_totals_from_excel(excel_path)


def replace_placeholders(document: Document, replacements: Dict[str, str]) -> None:
    def replace_in_paragraphs(paragraphs) -> None:
        for paragraph in paragraphs:
            text = paragraph.text
            for key, value in replacements.items():
                text = text.replace(key, value)
            paragraph.text = text

    replace_in_paragraphs(document.paragraphs)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_in_paragraphs(cell.paragraphs)


def generate_report(template_path: Path, output_path: Path, data: ExtractedData) -> None:
    if not template_path.exists():
        raise FileNotFoundError(f"Word template not found: {template_path}")

    doc = Document(template_path)
    replacements = {
        "{{court}}": data.court,
        "{{case_no}}": data.case_no,
        "{{car_model}}": data.car_model,
        "{{car_colour}}": data.car_colour,
        "{{plate_no}}": data.plate_no,
        "{{judge}}": data.judge,
        "{{judge_short}}": data.judge_short,
        "{{total_amount}}": data.total_amount,
        "{{total_amount_str}}": data.total_amount_str,
        "{{incident_date}}": data.accident_date,
        "{{report_date}}": data.report_date,
    }
    replace_placeholders(doc, replacements)
    doc.save(output_path)


def process_pdf(
    pdf_path: Path,
    excel_path: Path,
    template_path: Path,
    output_filename: Optional[str] = None,
) -> Path:
    text = read_pdf_text(pdf_path)
    pdf_data, missing_fields = extract_pdf_data(text)
    logging.info("PDF %s extracted fields: %s", pdf_path.name, pdf_data)
    if missing_fields:
        logging.warning("PDF %s missing fields: %s", pdf_path.name, ", ".join(missing_fields))

    total_amount = extract_total_amount_from_excel(excel_path)
    total_amount_str = amount_to_az_text(total_amount)
    if total_amount:
        logging.info("PDF %s extracted total_amount from Excel: %s", pdf_path.name, total_amount)
    else:
        logging.warning("PDF %s missing field: total_amount (Excel D/E in 'Yekun dəyər' row).", pdf_path.name)
    if total_amount_str:
        logging.info("PDF %s total_amount_str: %s", pdf_path.name, total_amount_str)
    else:
        logging.warning("PDF %s could not build total_amount_str from total_amount.", pdf_path.name)

    judge_value = pdf_data.get("judge", "")
    judge_short = get_judge_short_label(judge_value)

    extracted = ExtractedData(
        court=pdf_data.get("court", ""),
        case_no=pdf_data.get("case_no", ""),
        car_model=pdf_data.get("car_model", ""),
        car_colour=pdf_data.get("car_colour", DEFAULT_CAR_COLOUR),
        plate_no=pdf_data.get("plate_no", ""),
        judge=judge_value,
        judge_short=judge_short,
        accident_date=pdf_data.get("accident_date", ""),
        total_amount=total_amount,
        total_amount_str=total_amount_str,
        report_date=datetime.now().strftime("%d.%m.%Y"),
    )

    if output_filename:
        output_file = OUTPUT_DIR / output_filename
    else:
        output_file = OUTPUT_DIR / f"report_{pdf_path.stem}.docx"
    generate_report(template_path, output_file, extracted)
    return output_file


def main() -> None:
    ensure_directories()
    setup_logging()

    toyota_pdf = INPUT_PDF_DIR / "Qərardad 183.26.pdf"
    if toyota_pdf.exists():
        pdf_files = [toyota_pdf]
    else:
        pdf_files = sorted(INPUT_PDF_DIR.glob("*.pdf"))

    if not pdf_files:
        raise FileNotFoundError(
            f"No PDF files found in {INPUT_PDF_DIR}. Add files and rerun processor."
        )

    excel_path = find_first_file(TEMPLATES_DIR, ("*.xlsx", "*.xlsm", "*.xls"))
    if excel_path is None:
        raise FileNotFoundError(
            f"No Excel file found in {TEMPLATES_DIR}. Add damage calculation file."
        )

    toyota_template = TEMPLATES_DIR / "report_template.docx"
    if toyota_template.exists():
        template_path = toyota_template
    else:
        template_path = find_first_file(TEMPLATES_DIR, ("*.docx",))
    if template_path is None:
        raise FileNotFoundError(
            f"No Word template found in {TEMPLATES_DIR}. Add .docx template file."
        )

    for pdf_path in pdf_files:
        try:
            output_name = "Toyota_Report.docx" if pdf_path.name == "Qərardad 183.26.pdf" else None
            output_file = process_pdf(pdf_path, excel_path, template_path, output_name)
            print(f"[OK] Report created: {output_file}")
            logging.info("PDF %s processed successfully. Output: %s", pdf_path.name, output_file)
        except (FileNotFoundError, DataNotRecognizedError) as exc:
            print(f"[ERROR] {pdf_path.name}: {exc}")
            logging.error("PDF %s failed: %s", pdf_path.name, exc)
        except Exception as exc:  # noqa: BLE001
            print(f"[ERROR] {pdf_path.name}: Unexpected error: {exc}")
            logging.exception("PDF %s failed with unexpected error.", pdf_path.name)


if __name__ == "__main__":
    main()
